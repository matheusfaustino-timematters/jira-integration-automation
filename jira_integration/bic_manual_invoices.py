import csv
import io
import os
import subprocess
import time as time_sleep
from dataclasses import dataclass
from datetime import datetime
from datetime import time as dt_time
from difflib import get_close_matches
from pathlib import Path
from zoneinfo import ZoneInfo

import pandas as pd
from dotenv import load_dotenv
from jira import JIRA
from loguru import logger
from openpyxl import load_workbook
from server import Server

from jira_integration.types import JiraTransitionCodes

MANUAL_CREATION_TYPE = "manual"
ATTACHMENT_EXTENSIONS = (".xlsx", ".pdf")
STEERING_EXCEL_RELATIVE_PATH = "2_to_be_sent_to_AX/SPL_Invoices_for_AX_Sending.xlsx"
MANUAL_ROOT_RELATIVE_PATH = "1_created_manually"

REPORTS_PATH = "\\SAS Reports"
TASK_UPDATE_AX_SENDING_NAME = "AdHoc_SPL_BIC_2_Update_AX_Sending"
TASK_COPY_NAME = "AdHoc_SPL_BIC_5_Copy_Attachments"
TASK_SEND_NAME = "AdHoc_SPL_BIC_6_Send_Attachments"

CMD_TASK_RUN_STR = 'Start-ScheduledTask -TaskName "{task_name}" -TaskPath "{task_path}" | ConvertTo-Csv -NoTypeInformation'
CMD_TASK_STATUS_STR = 'Get-ScheduledTask | Where-Object {{ $_.TaskPath -like "{task_path}\\*" -and $_.TaskName -eq  "{task_name}" }} | ConvertTo-Csv -NoTypeInformation'

BERLIN_TZ = ZoneInfo("Europe/Berlin")
AX_PROCESSING_CUTOFF = dt_time(10, 0)

EMAIL_TO_DESIREE = "Desiree.Schaub@time-matters.com"
EMAIL_CC_TEMP = "matheus.faustino@time-matters.com"
NOTIFIED_MARKER = f"Notified {EMAIL_TO_DESIREE} about wrong file names"
SENT_MARKER = "Sent to AX: "


@dataclass
class ManualRow:
    invoice_no: str
    file_name: str
    folder: str


@dataclass
class MissingFile:
    row: ManualRow
    candidates: list[str]


@dataclass
class ManualFilesCheck:
    found: list[ManualRow]
    missing: list[MissingFile]

    @property
    def all_found(self) -> bool:
        return len(self.missing) == 0


def _attachments_root() -> Path:
    return Path(os.environ["SPL_BIC_ATTACHMENTS_ROOT"])


def is_manual_excel_up_to_date() -> bool:
    excel_path = _attachments_root() / STEERING_EXCEL_RELATIVE_PATH
    modified_date = datetime.fromtimestamp(
        excel_path.stat().st_mtime, tz=BERLIN_TZ
    ).date()
    return modified_date == datetime.now(BERLIN_TZ).date()


def _normalize_invoice_no(value) -> str:
    """Normalizes an InvoiceNo cell value so pandas (which upcasts a column with any
    blank/NaN entries to float64) and openpyxl (which preserves the original int/str
    type) always agree on the same string for the same invoice."""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def load_manual_rows() -> list[ManualRow]:
    excel_path = _attachments_root() / STEERING_EXCEL_RELATIVE_PATH
    logger.info(f"Reading manual rows from {excel_path}")

    df = pd.read_excel(excel_path)
    manual_df = df[df["creation_type"] == MANUAL_CREATION_TYPE]

    return [
        ManualRow(
            invoice_no=_normalize_invoice_no(row["InvoiceNo"]),
            file_name=str(row["file_name"]),
            folder=str(int(row["folder"])),
        )
        for _, row in manual_df.iterrows()
    ]


def check_manual_files(rows: list[ManualRow]) -> ManualFilesCheck:
    manual_root = _attachments_root() / MANUAL_ROOT_RELATIVE_PATH
    found: list[ManualRow] = []
    missing: list[MissingFile] = []

    for row in rows:
        folder_path = manual_root / row.folder
        file_found = any(
            (folder_path / f"{row.file_name}{ext}").exists()
            for ext in ATTACHMENT_EXTENSIONS
        )
        if file_found:
            logger.info(
                f"Found manual file for invoice {row.invoice_no}: {row.file_name}"
            )
            found.append(row)
            continue

        existing_names = (
            [
                p.stem
                for p in folder_path.iterdir()
                if p.suffix.lower() in ATTACHMENT_EXTENSIONS
            ]
            if folder_path.exists()
            else []
        )
        candidates = get_close_matches(row.file_name, existing_names, n=3, cutoff=0.6)
        logger.warning(
            f"Missing manual file for invoice {row.invoice_no}: {row.file_name} "
            f"(candidates: {candidates})"
        )
        missing.append(MissingFile(row=row, candidates=candidates))

    return ManualFilesCheck(found=found, missing=missing)


def run_scheduled_task_and_wait(server: Server, task_name: str) -> bool:
    logger.info(f"Running task '{task_name}'")
    server.run_ps_cmd(
        CMD_TASK_RUN_STR.format(task_name=task_name, task_path=REPORTS_PATH)
    )

    # assume running state because window cmd do not return anything when triggering the task on shell
    state = "Running"
    while state == "Running":
        logger.info(f"Check if task '{task_name}' still running")
        time_sleep.sleep(1)
        output_bytes = server.run_ps_cmd(
            CMD_TASK_STATUS_STR.format(task_name=task_name, task_path=REPORTS_PATH)
        )
        result = next(csv.DictReader(io.StringIO(output_bytes.decode("latin-1"))))
        state = result["State"]

    logger.info(f"Task '{task_name}' finished with status '{state}'")
    return state == "Ready"


def run_update_ax_sending(server: Server) -> bool:
    """Runs AdHoc_SPL_BIC_2, which refreshes the manual invoices Excel with new data."""
    return run_scheduled_task_and_wait(server, TASK_UPDATE_AX_SENDING_NAME)


def trigger_copy_and_send(server: Server) -> tuple[bool, str | None]:
    """Runs AdHoc_SPL_BIC_5 then, if it succeeds, AdHoc_SPL_BIC_6. Returns (success, failed_step)."""
    if not run_scheduled_task_and_wait(server, TASK_COPY_NAME):
        return False, TASK_COPY_NAME

    if not run_scheduled_task_and_wait(server, TASK_SEND_NAME):
        return False, TASK_SEND_NAME

    return True, None


def trigger_copy_and_send_for_rows(
    server: Server, rows_to_include: list[ManualRow]
) -> tuple[bool, str | None]:
    """Temporarily trims the steering Excel down to `rows_to_include` (plus all
    non-manual rows), runs AdHoc_SPL_BIC_5/6 against that trimmed file, then restores
    the original Excel so later checks still see every originally-expected row.

    Uses openpyxl to delete rows in place (rather than rebuilding the workbook from a
    pandas DataFrame) so other sheets and formatting survive the trim, and restores
    from a raw byte backup so the restore is exact regardless."""
    excel_path = _attachments_root() / STEERING_EXCEL_RELATIVE_PATH
    original_bytes = excel_path.read_bytes()
    include_invoices = {row.invoice_no for row in rows_to_include}

    workbook = load_workbook(excel_path)
    sheet = workbook.active
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    creation_type_col = header.index("creation_type")
    invoice_no_col = header.index("InvoiceNo")

    rows_to_delete = [
        row[0].row
        for row in sheet.iter_rows(min_row=2)
        if row[creation_type_col].value == MANUAL_CREATION_TYPE
        and _normalize_invoice_no(row[invoice_no_col].value) not in include_invoices
    ]

    logger.info(
        f"Trimming {excel_path}: keeping invoice(s) {sorted(include_invoices)}, "
        f"dropping {len(rows_to_delete)} manual row(s) for this AX run"
    )
    for row_idx in reversed(rows_to_delete):
        sheet.delete_rows(row_idx)

    try:
        workbook.save(excel_path)
        logger.info(f"Triggering {TASK_COPY_NAME} / {TASK_SEND_NAME} against trimmed Excel")
        result = trigger_copy_and_send(server)
        logger.info(f"Copy/send for trimmed Excel finished with result {result}")
        return result
    finally:
        excel_path.write_bytes(original_bytes)
        logger.info(f"Restored {excel_path} to its original contents")


def already_sent_invoices(
    jira: JIRA, issue_key: str, comments: list | None = None
) -> set[str]:
    sent: set[str] = set()
    for comment in comments if comments is not None else jira.comments(issue_key):
        for line in comment.body.splitlines():
            if SENT_MARKER in line:
                after_marker = line.split(SENT_MARKER, 1)[1]
                sent.update(
                    invoice_no.strip()
                    for invoice_no in after_marker.split(", ")
                    if invoice_no.strip()
                )
    logger.info(f"{issue_key}: {len(sent)} invoice(s) already sent in previous cycle(s): {sorted(sent)}")
    return sent


def run_copy_send_for_rows(
    jira: JIRA, issue_key: str, server: Server, rows_to_send: list[ManualRow]
) -> bool:
    invoice_list = ", ".join(row.invoice_no for row in rows_to_send)
    logger.info(f"{issue_key}: running copy/send for invoice(s) {invoice_list}")

    if is_past_ax_processing_cutoff():
        logger.warning(f"{issue_key}: running past the 10:00 AX cutoff")
        jira.add_comment(
            issue_key,
            ":robot: Processed after the 10:00 AX cutoff - this may collide with the "
            "other bot's run, please check manually.",
            is_internal=True,
        )

    is_success, failed_step = trigger_copy_and_send_for_rows(server, rows_to_send)

    if is_success:
        logger.info(f"{issue_key}: successfully sent invoice(s) {invoice_list} to AX")
        jira.add_comment(
            issue_key,
            f":robot: {SENT_MARKER}{invoice_list}",
            is_internal=True,
        )
    else:
        logger.error(f"{issue_key}: copy/send failed while running '{failed_step}'")
        jira.add_comment(
            issue_key,
            f":robot: BipBop failed while running '{failed_step}', check the SAS scheduler manually",
            is_internal=True,
        )

    return is_success


def resolve_ticket_complete(jira: JIRA, issue_key: str) -> None:
    logger.info(f"{issue_key}: all manual invoices sent, resolving")
    jira.add_comment(
        issue_key,
        ":robot: BipBop finished task without errors",
        is_internal=True,
    )
    jira.transition_issue(issue_key, JiraTransitionCodes.RESOLVE_THIS_ISSUE.value)


def run_manual_invoices_cycle(
    jira: JIRA, issue_key: str, server: Server, rows: list[ManualRow]
) -> bool:
    """Sends whatever manual invoices are currently found (skipping any already sent
    in a previous cycle), notifies Desiree about anything still missing, and only
    resolves the ticket once nothing is left missing."""
    check = check_manual_files(rows)
    comments = jira.comments(issue_key)
    already_sent = already_sent_invoices(jira, issue_key, comments)
    to_send = [row for row in check.found if row.invoice_no not in already_sent]

    if to_send:
        logger.info(f"{issue_key}: sending {len(to_send)} found manual invoice(s) to AX")
        if not run_copy_send_for_rows(jira, issue_key, server, to_send):
            return False

    if check.missing:
        logger.info(f"{issue_key}: {len(check.missing)} manual file(s) still missing")
        return _run_missing_files_unhappy_path(jira, issue_key, check.missing, comments)

    resolve_ticket_complete(jira, issue_key)
    return True


def is_past_ax_processing_cutoff() -> bool:
    return datetime.now(BERLIN_TZ).time() >= AX_PROCESSING_CUTOFF


def already_notified_desiree(
    jira: JIRA, issue_key: str, comments: list | None = None
) -> bool:
    return any(
        NOTIFIED_MARKER in c.body
        for c in (comments if comments is not None else jira.comments(issue_key))
    )


def format_missing_files_comment(missing: list[MissingFile]) -> str:
    lines = [":robot: Missing/misnamed manual invoice attachment(s):"]
    for m in missing:
        candidates = ", ".join(m.candidates) if m.candidates else "no close match found"
        lines.append(
            f"- Invoice {m.row.invoice_no} (folder {m.row.folder}): expected "
            f'"{m.row.file_name}" - possible candidate(s): {candidates}'
        )
    return "\n".join(lines)


def format_missing_files_email_body(missing: list[MissingFile]) -> str:
    lines = [
        "Hi Desiree,",
        "",
        "We found wrong file names in the manual invoice attachments you uploaded for AX sending.",
        "Please rename them to match exactly (the Excel file_name column is the source of truth):",
        "",
    ]
    for m in missing:
        current = (
            m.candidates[0] if m.candidates else "(no matching file found on disk)"
        )
        lines.append(f"- Folder {m.row.folder}, invoice {m.row.invoice_no}:")
        lines.append(f"    current file found: {current}")
        lines.append(f"    should be renamed to: {m.row.file_name}")
    lines.append("")
    lines.append("Please fix and let us know, we will retry sending automatically.")
    return "\n".join(lines)


def _run_missing_files_unhappy_path(
    jira: JIRA,
    issue_key: str,
    missing: list[MissingFile],
    comments: list | None = None,
) -> bool:
    jira.add_comment(
        issue_key,
        format_missing_files_comment(missing),
        is_internal=True,
    )

    if not already_notified_desiree(jira, issue_key, comments):
        sent = send_missing_files_email(
            subject=f"[{issue_key}] Wrong file names in manual AX invoice attachments",
            message=format_missing_files_email_body(missing),
        )
        if sent:
            jira.add_comment(issue_key, NOTIFIED_MARKER, is_internal=True)
        else:
            jira.add_comment(
                issue_key,
                ":robot: Failed to send notification email to Desiree, check manually",
                is_internal=True,
            )

    return True


def send_missing_files_email(subject: str, message: str) -> bool:
    load_dotenv(".env.local")

    mail_app = os.environ["MAILSEND_APP"]
    port = os.environ["MAILSEND_PORT"]
    smtp_server = os.environ["MAILSEND_SERVER"]
    account = os.environ["MAILSEND_ACCOUNT"]
    password = os.environ["MAILSEND_PASS"]

    cmd = [
        mail_app,
        "-starttls",
        "-port",
        port,
        "-auth",
        "-smtp",
        smtp_server,
        "-to",
        EMAIL_TO_DESIREE,
        "-cc",
        account,
        # temporary
        "-cc",
        "Mariia.Krasnopolska@time-matters.com",
        # "-cc",
        # EMAIL_CC_TEMP,
        "-from",
        account,
        "-sub",
        subject,
        "+bc",
        "-v",
        "-user",
        account,
        "-pass",
        password,
        "-M",
        message,
    ]

    result = subprocess.run(cmd, capture_output=True, text=True)
    if result.returncode != 0:
        logger.error(
            f"mailsend failed (code {result.returncode}): {result.stdout} {result.stderr}"
        )
        return False

    logger.info(f"Sent email to {EMAIL_TO_DESIREE}: {subject}")
    return True
